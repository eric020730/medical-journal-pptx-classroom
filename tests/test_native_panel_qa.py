import json
import sys
import tempfile
import unittest
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'tools'))
from qa_check import check_native_panels


class NativePanelTests(unittest.TestCase):
    def test_missing_duplicate_wrong_color_position_and_inventory(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            image=root/'figure.png'; Image.new('RGB',(100,100)).save(image)
            geom=[{'label':'A','fx_right':1,'fy_center':0.9}]
            image.with_suffix('.png.postprocess.json').write_text(json.dumps({'native_labels':True,'labels':['A'],'geometry':geom}))
            plan=root/'plan.json'
            plan.write_text(json.dumps({'assets':[{'id':'Figure_1','expected_labels':['A'],'export_image_panels':True}]}))
            spec={'meta':{'panel_crop_plan':str(plan)},'slides':[{'type':'figure','image':str(image),'caption':'Figure 1.'}]}
            prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(str(image),0,0,width=Inches(2),height=Inches(2))
            check=lambda:check_native_panels(prs,spec,root/'spec.json')
            self.assertTrue(check())
            box=slide.shapes.add_textbox(Inches(1.4),Inches(1.65),Inches(.55),Inches(.3))
            run=box.text_frame.paragraphs[0].add_run();run.text='A';run.font.size=Pt(18);run.font.color.rgb=RGBColor.from_string('8FA8C8')
            self.assertEqual(check(),[])
            run.font.color.rgb=RGBColor.from_string('000000');self.assertTrue(check())
            run.font.color.rgb=RGBColor.from_string('8FA8C8')
            box.left+=Inches(.1);self.assertTrue(check());box.left-=Inches(.1)
            duplicate=slide.shapes.add_textbox(0,0,Inches(1),Inches(1));duplicate.text='A'
            self.assertTrue(check())
            # Removing optional slide label fields cannot hide inventory errors.
            image.with_suffix('.png.postprocess.json').write_text('{}')
            self.assertTrue(check())
