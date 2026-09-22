"""Integrated asset regressions run in clean interpreters to avoid tools imports."""
import subprocess
import sys
import textwrap
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1] / '.agents/skills/medical-journal-to-pptx-integrated/scripts'
PRELUDE = r'''
import copy, hashlib, io, json, subprocess, sys, tempfile
from pathlib import Path
import pymupdf as fitz
from PIL import Image
import source_crops as crops
import image_polarity as polarity
import postprocess_assets as post
import vector_table
import qa_check
from extract_from_pdf import extract
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

temporary = tempfile.TemporaryDirectory()
root = Path(temporary.name)
pdf = root / 'source.pdf'
with fitz.open() as doc:
    page = doc.new_page(width=400, height=300)
    for y, text in [(30, 'TABLE 1'), (55, 'Header'), (100, 'First row'),
                    (150, 'Last row'), (180, 'Footnote')]:
        page.insert_text((20, y), text)
    page = doc.new_page(width=400, height=300)
    for x, label in [(20, 'A'), (160, 'B')]:
        im = Image.new('RGB', (100,100), 'white')
        for y in range(2,98):
            for z in range(2,98):
                v = 20 + (y * 3 + z * 7) % 170
                im.putpixel((z,y), (v,v,v))
        data = io.BytesIO(); im.save(data, format='PNG')
        page.insert_image(fitz.Rect(x,20,x+100,120), stream=data.getvalue())
        page.insert_text((x+85,135), label)
    doc.save(pdf)
table = {'id':'Table_1', 'type':'table', 'page':1, 'bbox':[10,10,280,190],
         'header_bottom':70, 'splits':[120],
         'expected_text':['TABLE 1','Header','Last row','Footnote']}
figure = {'id':'Figure_1', 'type':'figure', 'export_image_panels':True,
          'expected_labels':['A','B'], 'panels':[
              {'label':'A','page':2,'bbox':[10,10,130,145]},
              {'label':'B','page':2,'bbox':[150,10,270,145]}]}
plan = {'pdf':str(pdf),'source_sha256':crops.digest(pdf),'dpi':72,
        'expected_assets':['Table_1','Figure_1'],'assets':[table,figure]}
def meta(path):
    return json.loads(path.with_suffix(path.suffix+'.postprocess.json').read_text())
def save_meta(path, value):
    path.with_suffix(path.suffix+'.postprocess.json').write_text(json.dumps(value))
def audit():
    extracted = root/'extracted'
    extract(str(pdf), str(extracted), dpi=72, table_dpi=72)
    report = polarity.audit_extraction(extracted/'manifest.json', persist=False)
    assert report['ok'], report['failures']
    return report
def final(path, report):
    spec = root/'spec.json'
    spec.write_text(json.dumps({'slides':[{'type':'figure','image':str(path)}]}))
    return polarity.audit_final_assets(spec, report)
def reviewed_plan():
    # Simulate an explicit full-panel review of exact rendered pixels.
    with fitz.open(pdf) as doc:
        for panel in figure['panels']:
            image = crops.render(doc[1], crops.image_region(doc[1], panel['bbox']),72)
            panel['source_panel_label'] = {'status':'verified-absent','verified_absent':True,
                'absence_evidence':{'method':'full-panel-decoded-rgb-review-v1',
                    'review_box_px':[0,0,image.width,image.height],
                    'decoded_rgb_sha256':hashlib.sha256(image.tobytes()).hexdigest()}}
    return plan
def compose(out, *extra):
    result = subprocess.run([sys.executable, str(SCRIPTS/'recompose_panels_banded.py'),
        '--inputs', str(out/'Figure_1_A_image.png'),str(out/'Figure_1_B_image.png'),
        '--labels','A,B',str(out/'final.png'),
        '--geometry',str(out/'geometry.json'), *extra],capture_output=True,text=True)
    return result
'''


class IntegratedCropAssetTests(unittest.TestCase):
    def case(self, code):
        program = ('import sys\nfrom pathlib import Path\n'
                   f'SCRIPTS = Path({str(SCRIPTS)!r})\n'
                   'sys.path.insert(0, str(SCRIPTS))\n' + PRELUDE + '\n' + textwrap.dedent(code))
        result = subprocess.run([sys.executable, '-c', program], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_table_replay_and_final_safety_contract(self):
        self.case('''
            out=root/'out'; crops.generate(plan,out)
            report=audit()
            for name in ['Table_1A','Table_1B']:
                path=out/(name+'.png')
                assert not post.validate_final_sidecar(path,meta(path))
                assert not crops.replay_source_crop(path,meta(path))
                result=final(path,report); assert result['ok'],result
            assert post.validate_final_sidecar(out/'Figure_1.png',meta(out/'Figure_1.png'))
            assert post.validate_final_sidecar(out/'Figure_1_A_image.png',meta(out/'Figure_1_A_image.png'))
        ''')

    def test_tampered_pixels_hash_metadata_and_sources_fail_replay(self):
        self.case('''
            out=root/'out'; crops.generate(plan,out)
            path=out/'Table_1A.png'; original=meta(path)
            for key,value in [('output_id','Table_1B'),('source_sha256','0'*64),
                              ('dpi',144),('margin',0),('intermediate',True),
                              ('source',str(root/'missing.pdf'))]:
                changed=copy.deepcopy(original); changed[key]=value
                assert crops.replay_source_crop(path,changed), key
            with Image.open(path) as im:
                im=im.convert('RGB'); im.putpixel((30,30),(250,0,0)); im.save(path)
            changed=copy.deepcopy(original); changed['output_sha256']=crops.digest(path)
            assert crops.replay_source_crop(path,changed)
        ''')

    def test_wrong_audit_pdf_hash_page_and_failed_audit_rejected(self):
        self.case('''
            out=root/'out'; crops.generate(plan,out); report=audit()
            path=out/'Table_1A.png'
            for key,value in [('source_pdf_sha256','0'*64),('source_pdf',str(root/'other.pdf')),
                              ('ok',False),('verified_references',[])]:
                changed=copy.deepcopy(report); changed[key]=value
                assert not final(path,changed)['ok'],key
            # Declaring an arbitrary direct-PDF helper is still forbidden.
            changed=meta(path); changed['command']='same-width'; save_meta(path,changed)
            assert not final(path,report)['ok']
        ''')

    def test_image_only_rejects_whole_text_and_partial_glyph(self):
        self.case('''
            with fitz.open(pdf) as doc:
                page=doc[1]; page.insert_text((40,70),'C')
                for box in [[20,20,120,120],[20,20,43,120]]:
                    try: crops.region(page,box,image_only=True)
                    except ValueError: pass
                    else: raise AssertionError('image-only export accepted text')
            with fitz.open(pdf) as doc:
                # A font line box alone outside actual glyphs is still allowed.
                doc[1].insert_text((110,131),'D',fontsize=11)
                crops.image_region(doc[1],[10,10,130,145])
        ''')

    def test_invalid_inventory_anchors_and_page_do_not_write_outputs(self):
        self.case('''
            mutations=[lambda p:p['expected_assets'].append('Figure_1'),
                       lambda p:p['assets'][0].update(expected_text=['']),
                       lambda p:p['assets'][0].update(expected_text=[]),
                       lambda p:p['assets'][0].update(page=True),
                       lambda p:p.update(dpi=True),
                       lambda p:p['assets'][1].update(export_image_panels='yes')]
            for n, mutate in enumerate(mutations):
                bad=copy.deepcopy(plan); mutate(bad); out=root/str(n)
                try: crops.generate(bad,out)
                except ValueError: pass
                else: raise AssertionError(n)
                assert not out.exists()
        ''')

    def test_native_review_compositor_replay_and_edge_cleanup(self):
        self.case('''
            out=root/'out'; crops.generate(reviewed_plan(),out)
            result=compose(out,'--medical-image'); assert result.returncode==0,result.stderr
            sidecar=meta(out/'final.png')
            assert sidecar['native_labels'] is True, sidecar
            assert sidecar['safety_margin_px']==0
            assert not post.validate_final_sidecar(out/'final.png',sidecar)
            result=final(out/'final.png',audit()); assert result['ok'],result
            bad=meta(out/'Figure_1_A_image.png'); bad['image_region']['bbox'][0]=21
            save_meta(out/'Figure_1_A_image.png',bad)
            handled,errors=polarity._deterministic_helper_evidence(out/'Figure_1_A_image.png',bad)
            assert handled and errors
        ''')

    def test_unknown_and_stale_label_review_not_promoted_to_native(self):
        self.case('''
            out=root/'out'; crops.generate(plan,out)
            result=compose(out); assert result.returncode==0,result.stderr
            assert meta(out/'final.png')['native_labels'] is False
            reviewed_plan()
            figure['panels'][0]['source_panel_label']['absence_evidence']['decoded_rgb_sha256']='0'*64
            try: crops.generate(plan,root/'bad')
            except ValueError: pass
            else: raise AssertionError('stale review accepted')
            assert not (root/'bad').exists()
        ''')

    def test_safe_no_trim_preserved_and_residual_white_edges_fail_qa(self):
        self.case('''
            from recompose_panels_banded import residual_edge_review
            out=root/'out'; crops.generate(plan,out)
            for args in [('--medical-image','--no-trim'),
                         ('--asset-type','clinical-image','--no-trim')]:
                result=compose(out,*args)
                assert result.returncode==0,result.stderr
                sidecar=meta(out/'final.png')
                assert any(item['residual_edge_review']['status']=='needs-review'
                           for item in sidecar['panel_cleanup']),sidecar
                spec=root/'rim-spec.json'
                spec.write_text(json.dumps({'meta':{},'slides':[{'type':'figure',
                    'image':str(out/'final.png'),'caption':'Figure 1.'}]}))
                result=qa_check.validate_specification(spec,audit_images=False)
                assert any('residual frame is blocking' in error for error in result['failures']),result
            clean=Image.new('RGB',(100,100),(40,50,60))
            assert residual_edge_review(clean,include_bounded=True)['status']=='clear'
        ''')

    def test_vector_default_converter_matches_replay_despite_path_install(self):
        self.case('''
            from unittest import mock
            managed=root/'managed-soffice'; managed.touch()
            with mock.patch.object(vector_table.workflow,'find_binary',return_value=managed), \
                 mock.patch.object(vector_table.shutil,'which',return_value='/other/soffice'):
                assert vector_table._resolve_soffice()==managed
                assert vector_table._resolve_soffice('soffice')==managed
            explicit=root/'explicit-soffice';explicit.touch()
            assert vector_table._resolve_soffice(explicit)==explicit.resolve()
        ''')

    def test_vector_anchors_required_and_padding_excludes_neighbors(self):
        self.case('''
            for anchors in [[],[''],['absent']]:
                try: vector_table.canonical_svg(pdf,page=1,requested_bbox=[10,10,280,190],
                    pad_x=20,pad_top=20,pad_bottom=20,expected_text=anchors)
                except ValueError: pass
                else: raise AssertionError(anchors)
            # Neighbor inside the expanded padding extent must not appear in SVG.
            with fitz.open(pdf) as doc:
                doc[0].draw_rect(fitz.Rect(282,20,290,80),fill=(1,0,0))
                neighbor=root/'neighbor.pdf'; doc.save(neighbor)
            payload,bbox,aspect=vector_table.canonical_svg(neighbor,page=1,
                requested_bbox=[10,10,280,190],pad_x=20,pad_top=20,pad_bottom=20,
                expected_text=table['expected_text'])
            # PyMuPDF's SVG reader ignores clip paths. Verify with the actual
            # LibreOffice consumer, then rasterize its PDF with PyMuPDF.
            svg_path=root/'table.svg'; svg_path.write_bytes(payload)
            result=subprocess.run([str(vector_table._resolve_soffice()),
                '-env:UserInstallation='+ (root/'lo-profile').as_uri(),
                '--headless','--convert-to','pdf','--outdir',str(root),str(svg_path)],
                capture_output=True,text=True)
            assert result.returncode==0,result.stderr
            rendered=fitz.open(root/'table.pdf')
            pix=rendered[0].get_pixmap(alpha=False)
            im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
            assert any(max(pixel)<200 for pixel in im.getdata())
            assert not any(r>200 and g<80 and b<80 for r,g,b in im.getdata())
            assert bbox==[10,10,280,190] and abs(aspect-310/220)<1e-10
        ''')

    def test_native_inventory_binding_malformed_and_duplicate_checks(self):
        self.case('''
            out=root/'out'; crops.generate(reviewed_plan(),out)
            result=compose(out); assert result.returncode==0,result.stderr
            prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6])
            image=out/'final.png'
            picture=slide.shapes.add_picture(str(image),0,0,width=Inches(4),height=Inches(3))
            for geom in meta(image)['native_label_geometry']:
                right=geom['fx_right']*4-.05; center=geom['fy_center']*3
                box=slide.shapes.add_textbox(Inches(right-.55),Inches(center-.15),Inches(.55),Inches(.3))
                run=box.text_frame.paragraphs[0].add_run();run.text=geom['label']
                run.font.size=Pt(18);run.font.color.rgb=RGBColor.from_string('8FA8C8')
            spec={'meta':{'panel_crop_plan':str(out/'plan.json')},
                  'slides':[{'type':'figure','image':str(image),'source_asset_id':'figure:1'}]}
            def check(): return qa_check._check_source_inventory_panels(prs,spec,root/'spec.json',None)
            assert check()==[],check()
            original=meta(image)
            for field,value in [('source_inputs',list(reversed(original['source_inputs']))),
                                ('native_label_geometry','bad')]:
                changed=copy.deepcopy(original);changed[field]=value;save_meta(image,changed)
                assert check(),field
            save_meta(image,original)
            for mutation in [lambda p:p['assets'].append(copy.deepcopy(p['assets'][1])),
                             lambda p:p['assets'][1].update(expected_labels=[]),
                             lambda p:p['assets'][1].update(export_image_panels='yes')]:
                bad=copy.deepcopy(plan);mutation(bad);(out/'plan.json').write_text(json.dumps(bad))
                assert check()
        ''')

if __name__ == '__main__':
    unittest.main()
