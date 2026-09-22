"""Student entrypoint uses integrated gates and a project-pinned interpreter."""
from __future__ import annotations

import copy
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import classroom
import classroom_preflight as preflight


def style_report(style):
    return {"ok": True, "mode": "full", "style": style, "slides": 40,
            "image_polarity": True, "prebuild_qa": True, "postbuild_qa": True,
            "qa_receipt_current": True, "artifacts_verified": True,
            "render": {"pdf": "synthetic.pdf", "contact_sheet": "contact.jpg",
                       "preview_pages": 40}}


class IntegratedDispatchTests(unittest.TestCase):
    def test_student_commands_preserve_options_and_pin_project_interpreter(self):
        commands = [
            ["prepare", "paper.pdf", "--style", "nice", "--json"],
            ["init-run", "paper.pdf"],
            ["qa-spec", "spec.json", "--style", "nice", "--json"],
            ["build", "spec.json", "--out", "deck.pptx", "--style", "nice"],
            ["qa", "deck.pptx", "--spec", "spec.json", "--style", "nice"],
            ["render", "deck.pptx", "--preview", "--overwrite", "--json"],
            ["visual-review", "deck.pptx", "--evidence", "review.json", "--json"],
            ["qa-status", "deck.pptx", "--spec", "spec.json", "--require-delivery"],
            ["run", "source_crops", "--", "plan.json", "--out", "assets"],
            ["image-qa", "manifest.json", "--spec", "spec.json"],
        ]
        with patch.object(Path, "is_file", return_value=True), patch.object(classroom, "run_checked") as run:
            for args in commands:
                with self.subTest(command=args[0]):
                    original = list(args)
                    self.assertEqual(classroom.main(args), 0)
                    invoked = run.call_args.args[0]
                    self.assertEqual(invoked[:4], [str(classroom.PROJECT_PYTHON),
                        str(classroom.SKILL_SCRIPTS / "run.py"), "--runtime-python",
                        str(classroom.PROJECT_PYTHON)])
                    self.assertEqual(invoked[4:4 + len(args)], args)
                    self.assertEqual(args, original)
                    if args[0] in {"prepare", "init-run"}:
                        self.assertEqual(invoked[-4:], ["--workspace", str(ROOT),
                            "--output-dir", str(classroom.OUTPUTS_DIR)])

    def test_explicit_workspace_and_outputs_are_preserved_in_both_forms(self):
        cases = [
            (["--workspace", "custom workspace"], ["--output-dir", str(classroom.OUTPUTS_DIR)]),
            (["--workspace=custom workspace"], ["--output-dir", str(classroom.OUTPUTS_DIR)]),
            (["--output-dir", "custom outputs"], ["--workspace", str(ROOT)]),
            (["--output-dir=custom outputs"], ["--workspace", str(ROOT)]),
            (["--workspace", "custom workspace", "--output-dir", "custom outputs"], []),
            (["--workspace=custom workspace", "--output-dir=custom outputs"], []),
            (["--workspace=custom workspace", "--output-dir", "custom outputs"], []),
            (["--workspace", "custom workspace", "--output-dir=custom outputs"], []),
        ]
        with patch.object(Path, "is_file", return_value=True), patch.object(classroom, "run_checked") as run:
            for command in ("prepare", "init-run"):
                for options, defaults in cases:
                    with self.subTest(command=command, options=options):
                        args = [command, "paper.pdf", *options]
                        self.assertEqual(classroom.main(args), 0)
                        invoked = run.call_args.args[0]
                        self.assertEqual(invoked[:4], [str(classroom.PROJECT_PYTHON),
                            str(classroom.SKILL_SCRIPTS / "run.py"), "--runtime-python",
                            str(classroom.PROJECT_PYTHON)])
                        self.assertEqual(invoked[4:], [*args, *defaults])

    def test_missing_project_runtime_fails_without_global_fallback(self):
        with patch.object(Path, "is_file", return_value=False), patch.object(classroom, "run_checked") as run:
            with self.assertRaisesRegex(FileNotFoundError, "Project .venv"):
                classroom.main(["qa-spec", "spec.json"])
            run.assert_not_called()

    def test_global_overrides_and_python_injection_are_removed(self):
        with patch.dict(os.environ, {"MEDICAL_JOURNAL_PPTX_PYTHON": "global-python",
                "MEDICAL_JOURNAL_PPTX_RUNTIME": "global-runtime", "PYTHONHOME": "global-home",
                "PYTHONPATH": "global-modules"}), patch.object(classroom, "find_binary", return_value=None):
            env = classroom.subprocess_environment()
        self.assertEqual(env["MEDICAL_JOURNAL_PPTX_PYTHON"], str(classroom.PROJECT_PYTHON))
        self.assertEqual(env["MEDICAL_JOURNAL_PPTX_RUNTIME"], str(ROOT / ".venv"))
        self.assertNotIn("PYTHONHOME", env)
        self.assertNotIn("PYTHONPATH", env)
        self.assertEqual(env["PYTHONNOUSERSITE"], "1")

    def test_student_cannot_override_the_pinned_runtime(self):
        for argument in ("--runtime-python", "--runtime-python=global-python"):
            with self.subTest(argument=argument), patch.object(classroom, "run_checked") as run:
                with self.assertRaisesRegex(ValueError, "runtime is fixed"):
                    classroom.main(["qa-spec", "spec.json", argument, "global-python"])
                run.assert_not_called()

    def test_linked_venv_is_rejected(self):
        with patch.object(Path, "is_symlink", return_value=True):
            with self.assertRaisesRegex(RuntimeError, "symlink"):
                classroom.integrated_command(["qa-spec", "spec.json"])

    def test_helper_failure_is_not_swallowed(self):
        with patch.object(Path, "is_file", return_value=True), patch.object(classroom, "run_checked",
                side_effect=subprocess.CalledProcessError(7, ["qa-spec"])):
            with self.assertRaises(subprocess.CalledProcessError):
                classroom.main(["qa-spec", "bad.json"])


class IntegratedReadinessTests(unittest.TestCase):
    def test_smoke_runs_both_full_styles_through_integrated_runtime(self):
        responses = [subprocess.CompletedProcess([], 0, json.dumps(style_report(style)), "")
                     for style in ("standard", "nice")]
        with patch.object(Path, "is_file", return_value=True), patch.object(classroom, "ensure_directories"), \
                patch.object(classroom, "run_checked", side_effect=responses) as run, \
                patch.object(classroom, "verify_smoke_artifacts") as verify:
            report = classroom.smoke_test(keep=True, render=True)
        self.assertTrue(preflight.integrated_smoke_ready(report))
        self.assertEqual(verify.call_count, 2)
        for call, style in zip(run.call_args_list, ("standard", "nice")):
            args = call.args[0]
            self.assertIn("--runtime-python", args)
            self.assertEqual(args[args.index("--style") + 1], style)
            self.assertEqual(args[args.index("--mode") + 1], "full")
            self.assertIn("--render", args)
            self.assertIn("--keep", args)

    def test_partial_short_or_unrendered_smoke_never_attests_readiness(self):
        valid = {"styles": {style: style_report(style) for style in ("standard", "nice")}}
        broken = [{}, {"styles": {"standard": style_report("standard")}}, {"styles": []}]
        for style in ("standard", "nice"):
            for field, value in [("slides", 5), ("ok", False), ("prebuild_qa", False),
                                 ("postbuild_qa", False), ("image_polarity", False),
                                 ("qa_receipt_current", False), ("artifacts_verified", False),
                                 ("mode", "smoke"), ("render", {}), ("render", "bad")]:
                item = copy.deepcopy(valid)
                item["styles"][style][field] = value
                broken.append(item)
            item = copy.deepcopy(valid)
            item["styles"][style]["render"]["preview_pages"] = 39
            broken.append(item)
        for item in broken:
            with self.subTest(item=item):
                self.assertFalse(preflight.integrated_smoke_ready(item))

    def test_managed_windows_console_launcher_wins_over_path(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            executable = root / ".bootstrap/libreoffice/program/soffice.com"
            executable.parent.mkdir(parents=True)
            executable.touch()
            with patch.object(classroom, "PROJECT_ROOT", root), \
                    patch.object(classroom.platform, "system", return_value="Windows"), \
                    patch.object(classroom.shutil, "which", return_value="global-soffice.exe") as lookup:
                self.assertEqual(classroom.find_binary("soffice"), executable)
                lookup.assert_not_called()

    def test_actual_missing_preview_and_missing_receipt_block_readiness(self):
        from pptx import Presentation
        import pymupdf

        with tempfile.TemporaryDirectory() as temporary:
            work = Path(temporary)
            root = work / "integrated-smoke-fixture"
            root.mkdir()
            output = root / "synthetic_standard_full.pptx"
            deck = Presentation()
            for _ in range(40):
                deck.slides.add_slide(deck.slide_layouts[6])
            deck.save(output)
            pdf = root / "synthetic_standard_full.pdf"
            with pymupdf.open() as doc:
                for _ in range(40):
                    doc.new_page()
                doc.save(pdf)
            preview = root / "previews"
            preview.mkdir()
            contact = preview / "contact-sheet.jpg"
            contact.touch()
            report = style_report("standard")
            report.update(work_dir=str(root))
            report["render"].update(pdf=str(pdf), preview_dir=str(preview), contact_sheet=str(contact))
            with patch.object(classroom, "WORK_DIR", work), \
                    patch.object(classroom, "integrated_command", return_value=["qa-status"]), \
                    patch.object(classroom, "run_checked") as run:
                run.return_value = subprocess.CompletedProcess([], 0, json.dumps({"ok": True}))
                with self.assertRaisesRegex(RuntimeError, "render receipt"):
                    classroom.verify_smoke_artifacts(report, style="standard", render=True)
                run.return_value = subprocess.CompletedProcess([], 0, json.dumps({
                    "ok": True, "render": {"ok": True, "complete": True}}))
                with self.assertRaisesRegex(RuntimeError, "preview files"):
                    classroom.verify_smoke_artifacts(report, style="standard", render=True)
                run.return_value = subprocess.CompletedProcess([], 0, json.dumps({"ok": False}))
                with self.assertRaisesRegex(RuntimeError, "QA receipt"):
                    classroom.verify_smoke_artifacts(report, style="standard", render=True)

    def test_setup_and_launchers_use_integrated_tools_and_isolate_imports(self):
        for name in ("setup-codex.sh", "setup-codex.ps1"):
            text = (ROOT / name).read_text()
            self.assertIn("medical-journal-to-pptx-integrated", text)
            self.assertNotIn("medical-journal-to-pptx-classroom", text)
            self.assertIn("PYTHONNOUSERSITE", text)
        for name in ("journal", "journal.cmd"):
            text = (ROOT / name).read_text()
            self.assertIn("MEDICAL_JOURNAL_PPTX_PYTHON", text)
            self.assertIn("PYTHONNOUSERSITE", text)


if __name__ == "__main__":
    unittest.main()
