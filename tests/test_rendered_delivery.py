from __future__ import annotations

import importlib.util
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock

import pymupdf
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / '.agents/skills/medical-journal-to-pptx-integrated/scripts'
sys.path.insert(0, str(SCRIPTS))
import qa_attestation as qa
import render_attestation as receipts


def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


workflow = load('render_workflow_test', 'workflow.py')
runner = load('render_runner_test', 'run.py')


class DeliveryTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name).resolve()
        self.pptx = self.root / 'deck.pptx'
        deck = Presentation()
        for _ in range(2):
            deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(self.pptx)
        self.pdf = self.pptx.with_suffix('.pdf')
        with pymupdf.open() as pdf:
            for _ in range(2):
                pdf.new_page()
            pdf.save(self.pdf)
        self.spec = self.root / 'spec.json'
        self.spec.write_text('{}')
        self.previews = self.root / '.skill-work/previews/test'
        self.previews.mkdir(parents=True)
        for name in ('slide-1.jpg', 'slide-2.jpg', 'contact-sheet.jpg'):
            Image.new('RGB', (20, 20), 'white').save(self.previews / name)
        self.preview = {'preview_dir': str(self.previews), 'preview_pages': 2,
                        'contact_sheet': str(self.previews / 'contact-sheet.jpg')}
        self.evidence = self.root / 'review.json'
        self.automatic()

    def automatic(self):
        qa.record_qa(self.pptx, self.spec, mode='full', style='standard',
                     validate=lambda: {'ok': True, 'slides': 2, 'warnings': [], 'failures': []})

    def render(self, previews=True):
        return receipts.record_render(self.pptx, self.pdf,
            initial={'pptx_sha256': qa.digest(self.pptx), 'validator_sha256': qa.validator_digest()},
            pages=2, previews=self.preview if previews else {})

    def review(self, **overrides):
        evidence = {'render_receipt_sha256': qa.digest(receipts.receipt_path(self.pptx)),
                    'reviewer': 'Synthetic test reviewer', 'reviewed_pages': [1, 2],
                    'findings': 'Synthetic fixture pages checked for test only.', 'approved': True}
        evidence.update(overrides)
        self.evidence.write_text(json.dumps(evidence))
        return receipts.record_review(self.pptx, self.evidence)

    def status(self, strict=False):
        return qa.status(self.pptx, self.spec, require_delivery=strict)

    def test_structural_compatibility_and_explicit_full_delivery(self):
        self.assertTrue(self.status()['ok'])
        self.assertFalse(self.status(True)['ok'])
        self.render()
        self.assertTrue(self.status()['ok'])
        self.assertFalse(self.status()['visual_review']['ok'])
        self.assertFalse(self.status(True)['ok'])
        self.review()
        self.assertTrue(self.status(True)['delivery_ready'])
        self.spec.write_text('{"changed": true}')
        self.assertFalse(self.status(True)['ok'])

    def test_no_previews_cannot_claim_delivery(self):
        self.render(previews=False)
        self.assertTrue(self.status()['ok'])
        self.assertFalse(self.status(True)['ok'])
        with self.assertRaises(ValueError):
            self.review()

    def test_pdf_preview_contact_changes_and_deletions_are_stale(self):
        for path in (self.pdf, self.previews/'slide-1.jpg', self.previews/'contact-sheet.jpg'):
            original = path.read_bytes()
            for mutation in ('modify', 'delete'):
                with self.subTest(path=path.name, mutation=mutation):
                    path.write_bytes(original)
                    self.render()
                    if mutation == 'modify':
                        path.write_bytes(original + b'changed')
                    else:
                        path.unlink()
                    self.assertFalse(self.status()['ok'])
                    self.assertEqual(self.status()['status'], 'stale')
            path.write_bytes(original)

    def test_extra_preview_is_stale(self):
        self.render()
        (self.previews/'slide-3.jpg').write_bytes(b'extra')
        self.assertFalse(self.status()['ok'])

    def test_counts_and_missing_page_reject_receipt(self):
        self.preview['preview_pages'] = 1
        with self.assertRaises(RuntimeError):
            self.render()
        self.preview['preview_pages'] = 2
        (self.previews/'slide-2.jpg').rename(self.previews/'slide-3.jpg')
        with self.assertRaises(RuntimeError):
            self.render()

    def test_old_review_cannot_authorize_identical_rerender(self):
        self.render()
        self.review()
        self.render()
        self.assertFalse(self.status(True)['ok'])
        with self.assertRaisesRegex(ValueError, 'SHA-256'):
            receipts.record_review(self.pptx, self.evidence)

    def test_review_requires_exact_coverage_and_explicit_decision(self):
        for updates in ({'reviewed_pages': [1]}, {'reviewed_pages': [1, 1]},
                        {'reviewed_pages': [True, 2]}, {'approved': 'true'},
                        {'reviewer': ''}, {'findings': ''}, {'render_receipt_sha256': '0'*64}):
            with self.subTest(updates=updates):
                self.render()
                with self.assertRaises(ValueError):
                    self.review(**updates)
                self.assertFalse(self.status(True)['ok'])
        self.render()
        self.review(approved=False)
        self.assertFalse(self.status(True)['ok'])

    def test_failed_render_retry_invalidates_previous_success(self):
        self.render()
        self.review()
        with mock.patch.object(workflow, 'find_binary', return_value=None):
            with self.assertRaises(RuntimeError):
                workflow.render_presentation(self.pptx, overwrite=True, preview=True)
        self.assertFalse(self.status()['ok'])
        self.assertTrue(self.status()['automatic_qa']['ok'])

    def test_noop_converter_cannot_reuse_old_pdf(self):
        self.render()
        with mock.patch.object(workflow, 'find_binary', return_value=Path('soffice')), \
                mock.patch.object(workflow, 'run_checked'):
            with self.assertRaisesRegex(RuntimeError, 'fresh PDF'):
                workflow.render_presentation(self.pptx, overwrite=True)
        self.assertFalse(self.status()['ok'])

    def test_modified_pptx_during_conversion_is_rejected(self):
        original_pdf = self.pdf.read_bytes()
        def convert(command, **kwargs):
            (Path(command[command.index('--outdir')+1]) / self.pdf.name).write_bytes(original_pdf)
            self.pptx.write_bytes(self.pptx.read_bytes() + b'changed')
        with mock.patch.object(workflow, 'find_binary', return_value=Path('soffice')), \
                mock.patch.object(workflow, 'run_checked', side_effect=convert):
            with self.assertRaisesRegex(RuntimeError, 'changed during rendering'):
                workflow.render_presentation(self.pptx, overwrite=True)
        self.assertFalse(self.status()['ok'])

    def test_malformed_receipts_fail_closed_and_no_absolute_paths(self):
        self.render()
        path = receipts.receipt_path(self.pptx)
        text = path.read_text()
        self.assertNotIn(str(self.root), text)
        for malformed in ('[]', '{', '{"schema":true}'):
            path.write_text(malformed)
            self.assertFalse(self.status()['ok'])
        record = json.loads(text)
        record['pdf']['path'] = '../escape.pdf'
        path.write_text(json.dumps(record))
        self.assertFalse(self.status()['ok'])

    def test_real_render_has_two_previews_and_never_reviews(self):
        if not workflow.find_binary('soffice') or not workflow.find_binary('pdftoppm'):
            self.skipTest('LibreOffice and Poppler unavailable')
        result = workflow.render_presentation(self.pptx, overwrite=True, preview=True)
        self.assertEqual(result['preview_pages'], 2)
        self.assertFalse(result['visual_review'])
        self.assertFalse(result['delivery_ready'])
        self.assertTrue(self.status()['ok'])
        self.assertFalse(self.status(True)['ok'])


class RuntimeTests(unittest.TestCase):
    def test_explicit_runtime_precedes_global_and_environment_and_preserves_symlink(self):
        with tempfile.TemporaryDirectory() as directory:
            executable = Path(directory)/'python'
            executable.symlink_to(sys.executable)
            with mock.patch.dict(os.environ, {'MEDICAL_JOURNAL_PPTX_PYTHON': '/missing/global'}):
                self.assertEqual(runner.choose_python(str(executable)), executable)
            with self.assertRaises(FileNotFoundError):
                runner.choose_python(str(Path(directory)/'missing'))

    def test_managed_tool_wins_path_and_windows_uses_console(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            console = root/'.bootstrap/libreoffice/program/soffice.com'
            console.parent.mkdir(parents=True)
            console.touch()
            with mock.patch.dict(os.environ, {'MEDICAL_JOURNAL_PPTX_PROJECT_ROOT': str(root)}), \
                    mock.patch.object(workflow.platform, 'system', return_value='Windows'), \
                    mock.patch.object(workflow.shutil, 'which', return_value='/global/soffice.exe'):
                self.assertEqual(workflow.find_binary('soffice'), console.resolve())


if __name__ == '__main__':
    unittest.main()
